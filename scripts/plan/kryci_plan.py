#!/usr/bin/env python3
"""Denní krycí plán: měsíční rozpočet a cíl obratu rozpadnutý na dny a systémy. Jen čte.

Krycí plán říká, jak má měsíc vypadat den po dni: kolik který den v kterém systému utratit
a kolik má přijít obratu a konverzí. Každý den se s ním pak porovná skutečnost (plneni.py).

Použití:
    python3 scripts/plan/kryci_plan.py --klient muj-eshop --mesic 2026-11
    python3 scripts/plan/kryci_plan.py --klient muj-eshop --mesic 2026-11 --rozpocet 90000 --obrat 600000 \\
        --udalost 2026-11-27:+40% --udalost 2026-11-28..2026-11-30:+20%
    python3 scripts/plan/kryci_plan.py --klient ukazka --mesic 2026-10 --csv scripts/plan/ukazka-data.csv

Odkud se berou cíle (v tomhle pořadí):
1. --rozpocet / --obrat z příkazové řádky (celkem za všechny systémy),
2. `cile.mesicni_rozpocet` / `cile.mesicni_obrat` u účtů klienta v clients/hlidac-ucty.json,
3. realistický forecast (forecast.py) — když cíl nikdo nezadal.
Rozdělení mezi systémy: podle `cile.mesicni_rozpocet` v konfiguraci, jinak podle podílu útraty
za posledních 8 týdnů. Obrat systému = jeho rozpočet × jeho návratnost z historie.

Rozpad na dny: váhy dnů v týdnu z historie (16 týdnů, medián), státní svátky ČR (koeficient
z historie), ruční události --udalost (Black Friday, výprodej, výpadek). Součet dnů = měsíční cíl
na korunu.

Výstup: plans/<klient>/kryci-plan-<YYYY-MM>.csv (+ .xlsx, když je nainstalované openpyxl) + .md.
Návratový kód: 0 OK, 1 plán je, ale s varováním (krátká historie, chybí cíl), 2 chyba.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import data as D  # noqa: E402
import forecast as F  # noqa: E402
import model as M  # noqa: E402

ROOT = D.ROOT
SLOUPCE = ["datum", "den", "system", "plan_naklady", "plan_obrat", "plan_konverze", "plan_pno",
           "kum_plan_naklady", "kum_plan_obrat", "kum_plan_konverze", "poznamka"]


def _pomer(dny: dict, od: date, do: date, citatel: str, jmenovatel: str) -> float | None:
    j = M.soucet(dny, od, do, jmenovatel)
    return M.soucet(dny, od, do, citatel) / j if j > 0 else None


def sestav_plan(radky: list[M.Radek], rok: int, mesic: int, hist_do: date,
                rozpocet: float | None = None, obrat: float | None = None,
                cile: dict[str, dict] | None = None, udalosti: dict[date, float] | None = None) -> dict:
    """Čistý výpočet krycího plánu (bez souborů). Vrací cíle po systémech, řádky po dnech a popis metody."""
    cile = cile or {}
    udalosti = udalosti or {}
    radky = [r for r in radky if r.datum <= hist_do]
    varovani: list[str] = []
    zdroj: dict[str, str] = {}
    okno_od = hist_do - timedelta(days=55)

    systemy = sorted(set(M.systemy(radky)) | {s for s, c in cile.items() if c.get("mesicni_rozpocet")})
    if not systemy:
        raise SystemExit("Nemám data ani cíle po systémech — není z čeho plán sestavit.")

    # ---- rozpočet po systémech
    cfg_rozp = {s: float(cile[s]["mesicni_rozpocet"]) for s in systemy
                if cile.get(s, {}).get("mesicni_rozpocet")}
    fc = None

    def _forecast():
        nonlocal fc
        if fc is None:
            try:
                fc = F.spocti_forecast(radky, [(rok, mesic)], hist_do)["mesice"][0]["systemy"]
            except SystemExit as e:
                varovani.append(f"Forecast nejde: {e}")
                fc = {}
        return fc

    if cfg_rozp:
        podil = {s: cfg_rozp.get(s, 0.0) for s in systemy}
        zdroj["podil_systemu"] = "cile.mesicni_rozpocet v konfiguraci"
    else:
        podil = M.podily_systemu(radky, okno_od, hist_do, "naklady")
        podil = {s: podil.get(s, 0.0) for s in systemy}
        zdroj["podil_systemu"] = f"podíl útraty {okno_od:%d. %m.} – {hist_do:%d. %m. %Y} (8 týdnů)"
    if rozpocet:
        celkem_rozp = rozpocet
        zdroj["rozpocet"] = "zadaný --rozpocet"
    elif cfg_rozp:
        celkem_rozp = sum(cfg_rozp.values())
        zdroj["rozpocet"] = "cile.mesicni_rozpocet v konfiguraci"
    else:
        f = _forecast()
        if not f:
            raise SystemExit("Chybí rozpočet: zadej --rozpocet, nebo doplň cile.mesicni_rozpocet do konfigurace.")
        celkem_rozp = sum(x["realisticky"]["naklady"] for x in f.values())
        podil = {s: (f[s]["realisticky"]["naklady"] if s in f else 0.0) for s in systemy}
        zdroj["rozpocet"] = "realistický forecast (nikdo nezadal rozpočet — ověř s klientem)"
        zdroj["podil_systemu"] = "realistický forecast po systémech"
        varovani.append("Rozpočet nikdo nezadal, plán stojí na forecastu. Než ho pošleš klientovi, potvrď částku.")
    rozp_sys = M.rozdel_mezi_systemy(celkem_rozp, podil, 0)

    # ---- obrat po systémech: rozpočet × návratnost (ROAS) systému
    roas: dict[str, float | None] = {}
    for s in systemy:
        dny = M.po_dnech(radky, s)
        r_hist = _pomer(dny, okno_od, hist_do, "obrat", "naklady")
        f = _forecast().get(s) if (obrat is None and not any(c.get("mesicni_obrat") for c in cile.values())) else None
        if f and f["realisticky"]["naklady"] > 0:
            fr = f["realisticky"]
            x = (rozp_sys[s] / fr["naklady"] - 1) * 100 if fr["naklady"] else 0
            sc = M.scenar_rozpoctu(fr["naklady"], fr["obrat"], fr["konverze"], x, f["elasticita"])
            roas[s] = sc["obrat"] / rozp_sys[s] if rozp_sys[s] else r_hist
        else:
            roas[s] = r_hist
    cfg_obrat = {s: float(cile[s]["mesicni_obrat"]) for s in systemy if cile.get(s, {}).get("mesicni_obrat")}
    vaha_obratu = {s: rozp_sys[s] * (roas[s] or 0.0) for s in systemy}
    if obrat is not None:
        celkem_obrat = obrat
        zdroj["obrat"] = "zadaný --obrat; mezi systémy podle rozpočtu × návratnosti z historie"
        obrat_sys = M.rozdel_mezi_systemy(celkem_obrat, vaha_obratu if sum(vaha_obratu.values()) else podil, 0)
    elif cfg_obrat:
        obrat_sys = {s: cfg_obrat.get(s, round(vaha_obratu[s])) for s in systemy}
        zdroj["obrat"] = "cile.mesicni_obrat v konfiguraci (chybějící systémy dopočítané z návratnosti)"
    else:
        obrat_sys = {s: round(vaha_obratu[s]) for s in systemy}
        zdroj["obrat"] = ("dopočítaný: rozpočet × návratnost — z forecastu na tento měsíc (sezóna + klesající "
                          "výnos), kde forecast není, z posledních 8 týdnů")
        varovani.append("Cíl obratu nikdo nezadal, je dopočítaný z historie. Ověř ho s klientem.")
    for s in systemy:
        if rozp_sys[s] > 0 and roas[s] is None:
            varovani.append(f"{M.NAZVY_SYSTEMU.get(s, s)}: bez historie obratu — plán obratu je 0, doplň --obrat nebo cíl.")

    # ---- konverze: obrat / průměrná objednávka, u leadů útrata / CPA
    konv_sys: dict[str, float] = {}
    for s in systemy:
        dny = M.po_dnech(radky, s)
        aov = _pomer(dny, okno_od, hist_do, "obrat", "konverze")
        cpa = _pomer(dny, okno_od, hist_do, "naklady", "konverze")
        if obrat_sys[s] and aov:
            konv_sys[s] = round(obrat_sys[s] / aov, 1)
        elif cpa:
            konv_sys[s] = round(rozp_sys[s] / cpa, 1)
        else:
            konv_sys[s] = 0.0

    # ---- rozpad na dny
    radky_planu: list[dict] = []
    metoda: dict[str, dict] = {}
    for s in systemy:
        dny = M.po_dnech(radky, s)
        v_n, tydnu = M.vahy_dnu(dny, hist_do, "naklady")
        v_o, _ = M.vahy_dnu(dny, hist_do, "obrat")
        k_n, n_sv = M.koeficient_svatku(dny, "naklady")
        k_o, _ = M.koeficient_svatku(dny, "obrat")
        if tydnu < M.MIN_TYDNU:
            varovani.append(f"{M.NAZVY_SYSTEMU.get(s, s)}: jen {tydnu} celých týdnů historie — dny v týdnu mají "
                            f"rovnou váhu (nesuď z krátkého okna).")
        metoda[s] = {"vahy_naklady": v_n, "vahy_obrat": v_o, "tydnu": tydnu,
                     "svatek_naklady": k_n, "svatek_obrat": k_o, "svatku_v_historii": n_sv}
        n_dny = M.rozpad_mesice(rozp_sys[s], rok, mesic, v_n, k_n, udalosti, 0)
        o_dny = M.rozpad_mesice(obrat_sys[s], rok, mesic, v_o, k_o, udalosti, 0)
        k_dny = M.rozpad_mesice(konv_sys[s], rok, mesic, v_o, k_o, udalosti, 2)
        kn = ko = kk = 0.0
        for (d, n), (_, o), (_, k) in zip(n_dny, o_dny, k_dny):
            kn += n
            ko += o
            kk += k
            pozn = []
            sv = M.svatky_cr(d.year).get(d)
            if sv:
                pozn.append(sv)
            if d in udalosti:
                pozn.append(f"událost {(udalosti[d] - 1) * 100:+.0f} %")
            radky_planu.append({
                "datum": d.isoformat(), "den": M.DNY[d.weekday()], "system": s,
                "plan_naklady": round(n), "plan_obrat": round(o), "plan_konverze": round(k, 2),
                "plan_pno": round(n / o * 100, 2) if o else "",
                "kum_plan_naklady": round(kn), "kum_plan_obrat": round(ko), "kum_plan_konverze": round(kk, 2),
                "poznamka": "; ".join(pozn)})
    radky_planu.sort(key=lambda r: (r["datum"], r["system"]))
    return {"rok": rok, "mesic": mesic, "hist_do": hist_do, "systemy": systemy,
            "rozpocet": rozp_sys, "obrat": obrat_sys, "konverze": konv_sys, "radky": radky_planu,
            "metoda": metoda, "zdroj": zdroj, "udalosti": udalosti, "varovani": varovani}


# ---------------------------------------------------------------- výstupy
def uloz_csv(cesta: Path, plan: dict) -> None:
    cesta.parent.mkdir(parents=True, exist_ok=True)
    with cesta.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=SLOUPCE)
        w.writeheader()
        w.writerows(plan["radky"])


def uloz_xlsx(cesta: Path, plan: dict) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "Plán po dnech"
    ws.append(SLOUPCE)
    for c in ws[1]:
        c.font = Font(bold=True)
    vikend = PatternFill("solid", fgColor="F2F2F2")
    for r in plan["radky"]:
        ws.append([r[k] for k in SLOUPCE])
        if r["den"] in ("so", "ne") or r["poznamka"]:
            for c in ws[ws.max_row]:
                c.fill = vikend
    for col in "DEHI":
        for c in ws[col][1:]:
            c.number_format = "#,##0"
    ws.freeze_panes = "A2"
    s = wb.create_sheet("Souhrn")
    s.append(["system", "rozpocet", "obrat", "konverze", "pno_pct"])
    for c in s[1]:
        c.font = Font(bold=True)
    for sy in plan["systemy"]:
        o = plan["obrat"][sy]
        s.append([sy, plan["rozpocet"][sy], o, plan["konverze"][sy],
                  round(plan["rozpocet"][sy] / o * 100, 2) if o else None])
    n, o = sum(plan["rozpocet"].values()), sum(plan["obrat"].values())
    s.append(["celkem", n, o, round(sum(plan["konverze"].values()), 1), round(n / o * 100, 2) if o else None])
    wb.save(cesta)
    return True


def markdown(klient: str, plan: dict) -> str:
    rok, mesic = plan["rok"], plan["mesic"]
    n, o = sum(plan["rozpocet"].values()), sum(plan["obrat"].values())
    L = [f"# Krycí plán — {klient} — {F.nazev_mesice(rok, mesic)}", "",
         f"**Rozpočet {M.kc(n)}, obrat {M.kc(o)}, PNO {M.pct(M.pno(n, o))}, "
         f"konverze {M.cislo(sum(plan['konverze'].values()))}.**", "",
         "| Systém | Rozpočet | Obrat | Konverze | PNO | Průměr na den |", "|---|---:|---:|---:|---:|---:|"]
    dni = len(M.dny_mesice(rok, mesic))
    for s in plan["systemy"]:
        L.append(f"| {M.NAZVY_SYSTEMU.get(s, s)} | {M.kc(plan['rozpocet'][s])} | {M.kc(plan['obrat'][s])} | "
                 f"{M.cislo(plan['konverze'][s])} | {M.pct(M.pno(plan['rozpocet'][s], plan['obrat'][s]))} | "
                 f"{M.kc(plan['rozpocet'][s] / dni)} |")
    L += ["", "## Po dnech (všechny systémy)", "",
          "| Datum | Den | Náklady | Obrat | PNO | Kumulativ náklady | Kumulativ obrat | Poznámka |",
          "|---|---|---:|---:|---:|---:|---:|---|"]
    po_dnech: dict[str, dict] = {}
    for r in plan["radky"]:
        a = po_dnech.setdefault(r["datum"], {"den": r["den"], "n": 0, "o": 0, "pozn": r["poznamka"]})
        a["n"] += r["plan_naklady"]
        a["o"] += r["plan_obrat"]
    kn = ko = 0
    for d, a in po_dnech.items():
        kn += a["n"]
        ko += a["o"]
        L.append(f"| {date.fromisoformat(d):%d. %m.} | {a['den']} | {M.kc(a['n'])} | {M.kc(a['o'])} | "
                 f"{M.pct(M.pno(a['n'], a['o']))} | {M.kc(kn)} | {M.kc(ko)} | {a['pozn']} |")
    L += ["", "## Z čeho to vychází", "",
          f"- Rozpočet: {plan['zdroj']['rozpocet']}. Rozdělení mezi systémy: {plan['zdroj']['podil_systemu']}.",
          f"- Obrat: {plan['zdroj']['obrat']}.",
          f"- Historie pro váhy do {plan['hist_do']:%d. %m. %Y} (plán nevidí data z plánovaného měsíce)."]
    for s, m in plan["metoda"].items():
        vahy = " · ".join(f"{d} {M.cislo(v, 2)}" for d, v in zip(M.DNY, m["vahy_naklady"]))
        sv = (f"svátek = {M.cislo(m['svatek_naklady'], 2)}× běžného dne ({m['svatku_v_historii']} svátků v historii)"
              if m["svatku_v_historii"] >= 3 else
              f"svátek = 1,00× (v historii jen {m['svatku_v_historii']} svátků, koeficient nehádám)")
        L.append(f"- {M.NAZVY_SYSTEMU.get(s, s)} — váhy dnů útraty ({m['tydnu']} týdnů): {vahy}; {sv}.")
    if plan["udalosti"]:
        L.append("- Ruční události: " + ", ".join(f"{d:%d. %m.} {(x - 1) * 100:+.0f} %"
                                                  for d, x in sorted(plan["udalosti"].items())) +
                 ". Událost bere den větší díl na úkor ostatních dnů — měsíční součet se nemění.")
    L += ["", "Skutečnost proti plánu: `python3 scripts/plan/plneni.py --klient " + D.slug(klient) + "`.",
          "Plán je návrh. Denní rozpočty v účtech se mění jen po schválení (viz skill kryci-plan)."]
    if plan["varovani"]:
        L += ["", "## Upozornění", ""] + [f"- {x}" for x in plan["varovani"]]
    return "\n".join(L).rstrip() + "\n"


def main() -> int:
    ap = argparse.ArgumentParser(description="Denní krycí plán rozpočtu a obratu (jen čte)")
    ap.add_argument("--klient", required=True)
    ap.add_argument("--mesic", required=True, help="YYYY-MM")
    ap.add_argument("--rozpocet", type=float, help="měsíční rozpočet celkem (Kč)")
    ap.add_argument("--obrat", type=float, help="měsíční cíl obratu celkem (Kč)")
    ap.add_argument("--udalost", action="append", default=[],
                    help="ruční úprava dne, např. 2026-11-27:+40%% nebo 2026-12-21..2026-12-31:-50%% (opakovatelné)")
    ap.add_argument("--csv", type=Path, help="CSV export místo API (datum, system, naklady, obrat, konverze)")
    ap.add_argument("--ucty", type=Path, default=D.VYCHOZI_UCTY)
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()

    try:
        rok, mesic = M.parse_mesic(a.mesic)
        udalosti = M.sluc_udalosti(a.udalost)
        mimo = [d for d in udalosti if (d.year, d.month) != (rok, mesic)]
        if mimo:
            raise ValueError(f"Událost {mimo[0]} není v měsíci {a.mesic}.")
        zacatek = date(rok, mesic, 1)
        hist_do = min(date.today() - timedelta(days=1), zacatek - timedelta(days=1))
        od = date(hist_do.year - 2, hist_do.month, 1) - timedelta(days=31)
        radky, varovani = D.nacti_data(a.klient, a.csv, a.ucty, od, hist_do)
        konf = D.nacti_konfiguraci(a.ucty)
        cile = D.cile_systemu(D.ucty_klienta(konf, a.klient))
        if radky:
            hist_do = min(hist_do, max(r.datum for r in radky))
        plan = sestav_plan(radky, rok, mesic, hist_do, a.rozpocet, a.obrat, cile, udalosti)
        plan["varovani"] = varovani + plan["varovani"]
    except (SystemExit, ValueError) as e:
        print(f"CHYBA: {e}", file=sys.stderr)
        if a.json:
            print(json.dumps({"chyba": str(e)}, ensure_ascii=False))
        return 2

    slozka = ROOT / "plans" / D.slug(a.klient)
    nazev = f"kryci-plan-{rok}-{mesic:02d}"
    uloz_csv(slozka / f"{nazev}.csv", plan)
    xlsx = uloz_xlsx(slozka / f"{nazev}.xlsx", plan)
    md = markdown(a.klient, plan)
    (slozka / f"{nazev}.md").write_text(md, encoding="utf-8")
    if a.json:
        print(json.dumps({"mesic": f"{rok}-{mesic:02d}", "rozpocet": plan["rozpocet"], "obrat": plan["obrat"],
                          "konverze": plan["konverze"], "zdroj": plan["zdroj"], "varovani": plan["varovani"],
                          "radky": plan["radky"]}, ensure_ascii=False, indent=2))
    else:
        print(md)
    print(f"[uloženo] {(slozka / nazev).relative_to(ROOT)}.csv{' + .xlsx' if xlsx else ''} + .md"
          + ("" if xlsx else "  (XLSX vynecháno: chybí openpyxl)"), file=sys.stderr)
    return 1 if plan["varovani"] else 0


if __name__ == "__main__":
    sys.exit(main())
